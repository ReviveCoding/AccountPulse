"""Source integrity, manifest generation, and canonical source audit."""

from __future__ import annotations

import gzip
import hashlib
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook

from accountpulse.paths import REPO


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def csv_shape(path: Path) -> tuple[int, int]:
    if path.suffix.lower() == ".gz":
        with gzip.open(path, "rt", encoding="utf-8", newline="") as stream:
            header = stream.readline()
            rows = sum(1 for _ in stream)
        return rows, len(header.rstrip("\r\n").split(","))
    frame = pd.read_csv(path)
    return frame.shape


CRITEO_EXPECTED_SHA256 = "2716e1bf0fd157a93b5bf86924d9088419dfbac2022c6cd90030220634f616dc"
CRITEO_REVISION = "2424920019e49d52d72c13ac1143ec5d53af276b"


def validate_criteo_source(path: Path) -> str:
    """Fail closed unless the official v2.1 artifact has the preregistered digest."""
    if not path.is_file():
        raise FileNotFoundError(path)
    observed = sha256(path)
    if observed != CRITEO_EXPECTED_SHA256:
        raise ValueError(
            "SOURCE_INTEGRITY_FAILURE: Criteo v2.1 SHA-256 mismatch; "
            f"expected {CRITEO_EXPECTED_SHA256}, observed {observed}"
        )
    return observed


def build_manifest() -> dict[str, Any]:
    incoming = REPO / "data/incoming"
    sources: list[dict[str, Any]] = []
    definitions = [
        (
            "maven_crm_sales_opportunities",
            "Maven Analytics",
            "https://mavenanalytics.io/data-playground/crm-sales-opportunities",
            "complete official package acquired 2026-09-23",
            "user-authenticated official Maven download",
            "Maven Analytics usage terms; redistribution not assumed",
            "FICTITIOUS_PUBLIC_CRM_NOT_PRODUCTION",
            "Fictitious public CRM benchmark; not enterprise production evidence. "
            "The prior 499-row preview remains schema/audit evidence only.",
            sorted((incoming / "maven").glob("*.csv")),
        ),
        (
            "olist_marketing_and_ecommerce",
            "Olist via Kaggle",
            "olistbr/marketing-funnel-olist and olistbr/brazilian-ecommerce",
            "Kaggle current 2026-09-22",
            "authenticated Kaggle API",
            "CC-BY-NC-SA-4.0",
            "PUBLIC_REAL_SPARSE_FUNNEL",
            "First-contact predictors only; seller linkage is outcome-only; CC-BY-NC-SA-4.0.",
            sorted((incoming / "olist").glob("*.csv")) + sorted((incoming / "olist").glob("*.zip")),
        ),
        (
            "uci_online_retail_ii",
            "UCI Machine Learning Repository",
            "dataset 502 Online Retail II",
            "official archive current 2026-09-22",
            "direct UCI archive download",
            "CC BY 4.0",
            "PUBLIC_REAL_WHOLESALE_HEAVY_RETAIL",
            "Retail lifecycle benchmark, not CRM or enterprise production evidence; CC BY 4.0.",
            sorted((incoming / "uci").glob("*")),
        ),
        (
            "criteo_uplift_v2_1",
            "Criteo AI Lab via verified Criteo Hugging Face organization",
            "https://huggingface.co/datasets/criteo/criteo-uplift",
            f"v2.1; Hugging Face revision {CRITEO_REVISION}",
            "direct immutable-revision download from verified provider organization",
            "CC BY-NC-SA 4.0",
            "RANDOMIZED_CAUSAL_BENCHMARK_NOT_TRANSPORTABILITY_EVIDENCE",
            "Non-uniformly subsampled advertising experiment; absolute effects do not "
            "generalize beyond the released benchmark population.",
            sorted((incoming / "criteo").glob("*.csv.gz")),
        ),
    ]
    for (
        name,
        provider,
        identity,
        version,
        method,
        license_name,
        boundary,
        limitations,
        files,
    ) in definitions:
        records = []
        for path in files:
            if not path.is_file() or path.name == ".gitkeep":
                continue
            if name == "criteo_uplift_v2_1":
                validate_criteo_source(path)
            record: dict[str, Any] = {
                "filename": str(path.relative_to(REPO)),
                "size_bytes": path.stat().st_size,
                "sha256": sha256(path),
                "row_count": None,
                "column_count": None,
            }
            if path.suffix.lower() == ".csv" or path.name.lower().endswith(".csv.gz"):
                rows, columns = csv_shape(path)
                record.update(row_count=rows, column_count=columns)
            elif path.suffix.lower() == ".xlsx":
                workbook = load_workbook(path, read_only=True, data_only=True)
                record.update(
                    row_count=sum(sheet.max_row - 1 for sheet in workbook.worksheets),
                    column_count=max(sheet.max_column for sheet in workbook.worksheets),
                )
                workbook.close()
            records.append(record)
        time_range = None
        if name == "maven_crm_sales_opportunities":
            preview = pd.read_csv(incoming / "maven/sales_pipeline.csv")
            time_range = {
                "min": str(pd.to_datetime(preview["engage_date"]).min().date()),
                "max": str(pd.to_datetime(preview["engage_date"]).max().date()),
                "field": "engage_date",
            }
        elif name == "olist_marketing_and_ecommerce":
            leads = pd.read_csv(incoming / "olist/olist_marketing_qualified_leads_dataset.csv")
            time_range = {
                "min": str(pd.to_datetime(leads["first_contact_date"]).min().date()),
                "max": str(pd.to_datetime(leads["first_contact_date"]).max().date()),
                "field": "first_contact_date",
            }
        elif (
            name == "uci_online_retail_ii" and (REPO / "data/processed/uci_lines.parquet").exists()
        ):
            lines = pd.read_parquet(
                REPO / "data/processed/uci_lines.parquet", columns=["invoice_date"]
            )
            time_range = {
                "min": str(lines["invoice_date"].min()),
                "max": str(lines["invoice_date"].max()),
                "field": "InvoiceDate",
            }
        sources.append(
            {
                "source_name": name,
                "provider": provider,
                "official_source_identity": identity,
                "version": version,
                "license": license_name,
                "acquisition_method": method,
                "acquisition_timestamp": (
                    "2026-09-23T00:30:00Z"
                    if name == "criteo_uplift_v2_1"
                    else (
                        "2026-09-23T04:10:58.5605511Z"
                        if name == "maven_crm_sales_opportunities"
                        else "2026-09-22T22:33:00Z"
                    )
                ),
                "files": records,
                "time_range": time_range,
                "limitations": limitations,
                "claim_boundary": boundary,
            }
        )
    return {
        "schema_version": 1,
        "generated_at": datetime.now(UTC).isoformat(),
        "sources": sources,
    }


def write_manifest() -> None:
    manifest = build_manifest()
    (REPO / "SOURCE_MANIFEST.yaml").write_text(
        yaml.safe_dump(manifest, sort_keys=False), encoding="utf-8"
    )
    lines = [
        "# Data Sources",
        "",
        "Raw data is ignored by Git. Hashes and shapes are authoritative "
        "in `SOURCE_MANIFEST.yaml`.",
        "",
    ]
    for source in manifest["sources"]:
        lines.extend(
            [
                f"## {source['source_name']}",
                "",
                f"- Provider: {source['provider']}",
                f"- Official identity: {source['official_source_identity']}",
                f"- Version: {source['version']}",
                f"- License: {source['license']}",
                f"- Acquisition: {source['acquisition_method']}",
                f"- Claim boundary: `{source['claim_boundary']}`",
                f"- Limitations: {source['limitations']}",
                "",
            ]
        )
    (REPO / "DATA_SOURCES.md").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    write_manifest()
